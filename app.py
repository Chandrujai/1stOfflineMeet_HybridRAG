import os
import io
import re
import streamlit as st
import networkx as nx
import matplotlib.pyplot as plt
import pypdf
from docx import Document as DocxDocument
from dotenv import load_dotenv

# Load environment variables from .env
load_dotenv()

OPENAI_API_KEY = os.getenv("OPENAI_API_KEY")
if OPENAI_API_KEY:
    os.environ["OPENAI_API_KEY"] = OPENAI_API_KEY

# LangChain & NetworkX Imports
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_openai import OpenAIEmbeddings, ChatOpenAI
from langchain_chroma import Chroma
from langchain_experimental.graph_transformers import LLMGraphTransformer

# ==============================================================================
# Page Configuration
# ==============================================================================
st.set_page_config(
    page_title="Hybrid RAG: Vector DB + NetworkX Graph",
    page_icon="🕸️",
    layout="wide"
)

# ==============================================================================
# Session State Initialization
# ==============================================================================
if "messages" not in st.session_state:
    st.session_state.messages = []
if "vector_db" not in st.session_state:
    st.session_state.vector_db = None
if "graph" not in st.session_state:
    st.session_state.graph = nx.DiGraph()
if "indexed" not in st.session_state:
    st.session_state.indexed = False

# ==============================================================================
# Helper Functions: Robust File Parsers
# ==============================================================================
def parse_uploaded_file(uploaded_file):
    temp_filename = f"temp_{uploaded_file.name}"
    with open(temp_filename, "wb") as f:
        f.write(uploaded_file.getbuffer())
    
    docs = []
    file_ext = uploaded_file.name.split(".")[-1].lower()
    
    try:
        if file_ext == "pdf":
            reader = pypdf.PdfReader(temp_filename)
            text = ""
            for i, page in enumerate(reader.pages):
                page_text = page.extract_text()
                if page_text:
                    text += f"\n--- Page {i+1} ---\n" + page_text
            if text.strip():
                docs.append(Document(page_content=text, metadata={"source": uploaded_file.name}))
                
        elif file_ext in ["docx", "doc"]:
            doc = DocxDocument(temp_filename)
            text = "\n".join([p.text for p in doc.paragraphs if p.text.strip()])
            if text.strip():
                docs.append(Document(page_content=text, metadata={"source": uploaded_file.name}))
                
        elif file_ext in ["txt", "md"]:
            with open(temp_filename, "r", encoding="utf-8", errors="ignore") as f:
                text = f.read()
            if text.strip():
                docs.append(Document(page_content=text, metadata={"source": uploaded_file.name}))
    finally:
        if os.path.exists(temp_filename):
            os.remove(temp_filename)
            
    return docs

# ==============================================================================
# Helper Functions: NetworkX Graph Queries
# ==============================================================================
def query_networkx_graph(G, query_text):
    """Searches NetworkX graph for edges matching keywords in the query."""
    keywords = re.findall(r'\b[A-Z][a-zA-Z0-9_]{2,}\b', query_text)
    if not keywords:
        keywords = [w for w in query_text.split() if len(w) > 3]

    matching_relations = []
    for u, v, data in G.edges(data=True):
        relation_type = data.get("relation", "RELATED")
        if any(kw.lower() in u.lower() or kw.lower() in v.lower() for kw in keywords):
            matching_relations.append(f"({u}) -[{relation_type}]-> ({v})")
            
    return matching_relations[:25]

# ==============================================================================
# UI Layout: Sidebar Configuration
# ==============================================================================
st.title("🧠 Hybrid RAG Studio (Vector DB + NetworkX Graph)")

with st.sidebar:
    st.header("System Status")
    st.write(f"- **OpenAI API Key:** {'✅ Loaded' if OPENAI_API_KEY else '❌ Missing'}")
    st.write(f"- **Graph Engine:** `NetworkX (In-Memory)`")

    st.divider()
    st.header("Document Ingestion")
    uploaded_files = st.file_uploader(
        "Upload Documents (.pdf, .docx, .txt, .md)",
        type=["pdf", "docx", "doc", "txt", "md"],
        accept_multiple_files=True
    )
    
    chunk_size = st.slider("Chunk Size", min_value=200, max_value=2000, value=800, step=100)
    chunk_overlap = st.slider("Chunk Overlap", min_value=0, max_value=400, value=100, step=20)
    
    process_button = st.button("🚀 Process & Index Documents")

# ==============================================================================
# Document Ingestion Pipeline Orchestration
# ==============================================================================
if process_button:
    if not OPENAI_API_KEY:
        st.error("OpenAI API Key is missing. Check your environment or .env file.")
    elif not uploaded_files:
        st.warning("Please upload at least one document.")
    else:
        with st.spinner("Parsing documents & building Hybrid Vector + NetworkX Graph Indices..."):
            try:
                all_docs = []
                for file in uploaded_files:
                    docs = parse_uploaded_file(file)
                    all_docs.extend(docs)
                
                if not all_docs:
                    st.error("No extractable text found in the uploaded files.")
                else:
                    text_splitter = RecursiveCharacterTextSplitter(chunk_size=chunk_size, chunk_overlap=chunk_overlap)
                    chunks = text_splitter.split_documents(all_docs)
                    
                    # 1. Vector DB Pipeline (Chroma)
                    embeddings = OpenAIEmbeddings(model="text-embedding-3-small")
                    vector_db = Chroma.from_documents(
                        documents=chunks,
                        embedding=embeddings,
                        persist_directory="./chroma_db_store"
                    )
                    st.session_state.vector_db = vector_db
                    
                    # 2. NetworkX Knowledge Graph Pipeline via LLM Transformer
                    llm = ChatOpenAI(model="gpt-4o-mini", temperature=0)
                    transformer = LLMGraphTransformer(llm=llm)
                    graph_docs = transformer.convert_to_graph_documents(chunks)
                    
                    # Populate NetworkX DiGraph using .relationships instead of .rels
                    G = nx.DiGraph()
                    for gdoc in graph_docs:
                        for node in gdoc.nodes:
                            G.add_node(node.id, type=node.type)
                        for rel in gdoc.relationships:
                            G.add_edge(rel.source.id, rel.target.id, relation=rel.type)
                            
                    st.session_state.graph = G
                    st.session_state.indexed = True
                    st.success(f"Processing Complete! Created {len(chunks)} vector chunks & {G.number_of_nodes()} graph nodes.")
                
            except Exception as e:
                st.error(f"Error during ingestion pipeline: {e}")

# ==============================================================================
# UI Navigation Tabs
# ==============================================================================
tab_chat, tab_graph, tab_info = st.tabs(["💬 Hybrid RAG Chat", "🕸️ Interactive Graph Explorer", "ℹ️ System Info"])

with tab_chat:
    st.markdown("### Ask Questions Across Vector Context and Knowledge Graph")
    
    for message in st.session_state.messages:
        with st.chat_message(message["role"]):
            st.markdown(message["content"])

    user_query = st.chat_input("Ask a question about your documents...")
    if user_query:
        st.session_state.messages.append({"role": "user", "content": user_query})
        with st.chat_message("user"):
            st.markdown(user_query)

        if not st.session_state.get("indexed", False):
            response = "⚠️ Please upload and index documents using the sidebar first."
            st.session_state.messages.append({"role": "assistant", "content": response})
            with st.chat_message("assistant"):
                st.markdown(response)
        else:
            with st.chat_message("assistant"):
                with st.spinner("Retrieving semantic vector & structural graph knowledge..."):
                    # 1. Vector Search Retrieval
                    vector_results = st.session_state.vector_db.similarity_search(user_query, k=3)
                    vector_context = "\n\n".join([f"Chunk {i+1}:\n" + doc.page_content for i, doc in enumerate(vector_results)])

                    # 2. NetworkX Graph Traversal
                    graph_tuples = query_networkx_graph(st.session_state.graph, user_query)
                    graph_context = "\n".join(graph_tuples) if graph_tuples else "No direct entity relationships matched in the NetworkX graph."

                    # 3. Hybrid Context Synthesis using OpenAI
                    hybrid_prompt = f"""
You are an advanced AI assistant utilizing a Hybrid RAG framework. Synthesize a comprehensive answer using BOTH the vector similarity chunks and structural Knowledge Graph triplets provided below.

### 1. Vector Semantic Context (Document Chunks):
{vector_context}

### 2. Knowledge Graph Context (Extracted Entity Relationships):
{graph_context}

### User Query:
{user_query}

### Formulated Answer:
"""
                    llm = ChatOpenAI(model="gpt-4o-mini", temperature=0.2)
                    response_obj = llm.invoke(hybrid_prompt)
                    final_answer = response_obj.content

                    st.markdown(final_answer)

                    with st.expander("🔍 View Retrieved Hybrid Context Details"):
                        st.subheader("Semantic Vector Chunks")
                        st.text(vector_context)
                        st.subheader("Knowledge Graph Triplets")
                        st.text(graph_context)

                    st.session_state.messages.append({"role": "assistant", "content": final_answer})

with tab_graph:
    st.markdown("### NetworkX Knowledge Graph Visualizer")
    
    if st.session_state.get("indexed", False) and len(st.session_state.graph.nodes) > 0:
        if st.button("🔄 Refresh Graph View"):
            st.rerun()
            
        G = st.session_state.graph
        fig, ax = plt.subplots(figsize=(12, 7))
        pos = nx.spring_layout(G, k=0.4, seed=42)
        
        # Draw NetworkX Nodes & Edges
        nx.draw_networkx_nodes(G, pos, node_size=800, node_color="#4CAF50", alpha=0.9, ax=ax)
        nx.draw_networkx_edges(G, pos, width=1.5, alpha=0.5, edge_color="gray", ax=ax, arrows=True)
        nx.draw_networkx_labels(G, pos, font_size=8, font_weight="bold", ax=ax)
        
        edge_labels = nx.get_edge_attributes(G, 'relation')
        nx.draw_networkx_edge_labels(G, pos, edge_labels=edge_labels, font_size=6, ax=ax)
        
        plt.title("Extracted In-Memory Entities and Relationships", fontsize=12)
        plt.axis("off")
        st.pyplot(fig)
    else:
        st.info("Please complete the document ingestion step in the sidebar to render the graph visualization.")

with tab_info:
    st.markdown("### System Configuration Status")
    st.write(f"- **OpenAI API Key Configured:** {'✅ Yes' if OPENAI_API_KEY else '❌ No'}")
    st.write(f"- **Graph Backend:** NetworkX (In-Memory)")
    st.write(f"- **Documents Indexed:** {'✅ Yes' if st.session_state.get('indexed') else '❌ No'}")
    if st.session_state.get("indexed"):
        st.write(f"- **Total Graph Nodes:** {st.session_state.graph.number_of_nodes()}")
        st.write(f"- **Total Graph Edges:** {st.session_state.graph.number_of_edges()}")